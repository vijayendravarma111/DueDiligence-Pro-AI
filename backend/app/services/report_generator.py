import os
import json
import logging
from sqlalchemy.orm import Session
from docx import Document as DocxDocument
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.graphics.shapes import Drawing, Rect, String as DString, Line
from app.core.config import settings
from app.models.models import Report, Document, Company

logger = logging.getLogger(__name__)

# Colors matching professional dashboard
PRIMARY_COLOR = colors.HexColor("#0F172A")    # Dark slate
SECONDARY_COLOR = colors.HexColor("#1E293B")  # Muted panel
ACCENT_BLUE = colors.HexColor("#3B82F6")      # Electric blue
TEXT_DARK = colors.HexColor("#1E293B")
TEXT_MUTED = colors.HexColor("#64748B")
BG_LIGHT = colors.HexColor("#F8FAFC")

def draw_score_bar(score: int) -> Drawing:
    """Draws a clean score indicator bar (0-100) inside ReportLab."""
    d = Drawing(400, 30)
    # Background track
    d.add(Rect(0, 10, 400, 12, fillColor=colors.HexColor("#E2E8F0"), strokeColor=None))
    # Filled track
    fill_color = colors.HexColor("#EF4444") if score > 70 else (colors.HexColor("#F59E0B") if score > 40 else colors.HexColor("#10B981"))
    d.add(Rect(0, 10, int(score * 4), 12, fillColor=fill_color, strokeColor=None))
    # Score Text
    d.add(DString(410, 10, f"{score}/100", fontSize=11, fontName="Helvetica-Bold", fillColor=fill_color))
    return d

def build_pdf_report(report_id: int, db: Session) -> str:
    report = db.query(Report).filter(Report.id == report_id).first()
    if not report:
        raise ValueError("Report not found")
        
    doc_meta = db.query(Document).filter(Document.id == report.document_id).first()
    company = db.query(Company).filter(Company.id == report.company_id).first()
    
    company_name = company.name if company else "N/A"
    doc_name = doc_meta.file_name if doc_meta else "N/A"
    
    pdf_filename = f"report_{report_id}_{report.report_type}.pdf"
    pdf_path = os.path.join(settings.REPORT_DIR, pdf_filename)
    
    # Setup document
    doc = SimpleDocTemplate(
        pdf_path,
        pagesize=letter,
        rightMargin=54, leftMargin=54, topMargin=54, bottomMargin=54
    )
    
    styles = getSampleStyleSheet()
    
    # Custom Styles
    title_style = ParagraphStyle(
        'CoverTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=28,
        leading=34,
        textColor=PRIMARY_COLOR,
        alignment=TA_LEFT,
        spaceAfter=15
    )
    subtitle_style = ParagraphStyle(
        'CoverSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=14,
        leading=18,
        textColor=TEXT_MUTED,
        spaceAfter=50
    )
    meta_style = ParagraphStyle(
        'MetaText',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=14,
        textColor=TEXT_MUTED
    )
    h1_style = ParagraphStyle(
        'Header1',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=18,
        leading=22,
        textColor=PRIMARY_COLOR,
        spaceBefore=15,
        spaceAfter=10,
        keepWithNext=True
    )
    h2_style = ParagraphStyle(
        'Header2',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=16,
        textColor=ACCENT_BLUE,
        spaceBefore=10,
        spaceAfter=6,
        keepWithNext=True
    )
    body_style = ParagraphStyle(
        'BodyTextCustom',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=14,
        textColor=TEXT_DARK,
        spaceAfter=8
    )
    bullet_style = ParagraphStyle(
        'BulletCustom',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=14,
        textColor=TEXT_DARK,
        leftIndent=15,
        firstLineIndent=-10,
        spaceAfter=5
    )

    story = []
    
    # --- PAGE 1: COVER PAGE ---
    story.append(Spacer(1, 40))
    # Accent Bar
    d_line = Drawing(504, 6)
    d_line.add(Rect(0, 0, 504, 6, fillColor=ACCENT_BLUE, strokeColor=None))
    story.append(d_line)
    story.append(Spacer(1, 25))
    
    # Title & Subtitle
    story.append(Paragraph(report.title, title_style))
    story.append(Paragraph(f"AI-Powered Executive Intelligence Platform", subtitle_style))
    story.append(Spacer(1, 100))
    
    # Metadata Table
    meta_data = [
        [Paragraph(f"<b>COMPANY:</b> {company_name}", meta_style), Paragraph(f"<b>REPORT TYPE:</b> {report.report_type.upper()}", meta_style)],
        [Paragraph(f"<b>DOCUMENT:</b> {doc_name}", meta_style), Paragraph(f"<b>GENERATED:</b> {report.created_at.strftime('%Y-%m-%d %H:%M')}", meta_style)]
    ]
    t_meta = Table(meta_data, colWidths=[250, 250])
    t_meta.setStyle(TableStyle([
        ('LINEBELOW', (0,0), (-1,-1), 0.5, colors.HexColor("#E2E8F0")),
        ('TOPPADDING', (0,0), (-1,-1), 8),
        ('BOTTOMPADDING', (0,0), (-1,-1), 8),
    ]))
    story.append(t_meta)
    
    story.append(PageBreak())
    
    # --- PAGE 2: REPORT BODY ---
    data_content = json.loads(report.content)
    
    if report.report_type == "risk":
        story.append(Paragraph("Executive Summary", h1_style))
        story.append(Paragraph(data_content.get("executive_summary", ""), body_style))
        story.append(Spacer(1, 15))
        
        story.append(Paragraph("Risk Level Assessment", h1_style))
        overall_score = data_content.get("overall_risk_score", 0)
        story.append(draw_score_bar(overall_score))
        story.append(Paragraph(f"<b>Overall Risk Class:</b> {data_content.get('risk_level', 'Unknown')}", body_style))
        story.append(Spacer(1, 15))
        
        # Risk Breakdown Table
        story.append(Paragraph("Risk Category Breakdown", h2_style))
        breakdown = data_content.get("risk_breakdown", {})
        table_data = [["Category", "Score"]]
        for k, v in breakdown.items():
            name = k.replace("_", " ").title()
            table_data.append([name, f"{v}/100"])
        
        t_breakdown = Table(table_data, colWidths=[250, 100])
        t_breakdown.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), SECONDARY_COLOR),
            ('TEXTCOLOR', (0,0), (-1,0), colors.whitesmoke),
            ('ALIGN', (0,0), (-1,-1), 'LEFT'),
            ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
            ('FONTSIZE', (0,0), (-1,0), 10),
            ('BOTTOMPADDING', (0,0), (-1,-1), 6),
            ('TOPPADDING', (0,0), (-1,-1), 6),
            ('BACKGROUND', (0,1), (-1,-1), BG_LIGHT),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#CBD5E1")),
        ]))
        story.append(t_breakdown)
        story.append(Spacer(1, 15))
        
        story.append(Paragraph("Key Risks Identified", h1_style))
        for risk in data_content.get("key_risks", []):
            story.append(Paragraph(f"• {risk}", bullet_style))
        story.append(Spacer(1, 15))
        
        story.append(Paragraph("Mitigation Recommendations", h1_style))
        for rec in data_content.get("recommendations", []):
            story.append(Paragraph(f"• {rec}", bullet_style))
        story.append(Spacer(1, 15))
        
        story.append(Paragraph("Final Consultant Opinion", h1_style))
        story.append(Paragraph(data_content.get("final_consultant_opinion", ""), body_style))
        
    elif report.report_type == "investment":
        story.append(Paragraph("Investment Summary", h1_style))
        score = data_content.get("investment_score", 0)
        story.append(draw_score_bar(score))
        story.append(Paragraph(f"<b>Recommendation:</b> {data_content.get('recommendation', 'HOLD')}", body_style))
        story.append(Paragraph(f"<b>Analyst Confidence Score:</b> {data_content.get('confidence_score', 0)}%", body_style))
        story.append(Spacer(1, 15))
        
        story.append(Paragraph("Strategic Strengths", h1_style))
        for strength in data_content.get("strengths", []):
            story.append(Paragraph(f"• {strength}", bullet_style))
        story.append(Spacer(1, 15))
        
        story.append(Paragraph("Risk & Weaknesses", h1_style))
        for weakness in data_content.get("weaknesses", []):
            story.append(Paragraph(f"• {weakness}", bullet_style))
        story.append(Spacer(1, 15))
        
        story.append(Paragraph("Valuation Insights", h1_style))
        story.append(Paragraph(data_content.get("valuation_insight", ""), body_style))
        story.append(Spacer(1, 15))
        
        story.append(Paragraph("Final Analyst Opinion", h1_style))
        story.append(Paragraph(data_content.get("final_analyst_opinion", ""), body_style))
        
    else:  # Chat report
        story.append(Paragraph("AI Document Intelligence Query", h1_style))
        story.append(Paragraph(f"<b>Question:</b> {data_content.get('question', '')}", h2_style))
        story.append(Spacer(1, 10))
        story.append(Paragraph("Answer", h2_style))
        story.append(Paragraph(data_content.get("answer", ""), body_style))
        story.append(Spacer(1, 15))
        
        story.append(Paragraph("Supporting Document Evidence", h1_style))
        story.append(Paragraph(f"<i>\"{data_content.get('supporting_evidence', '')}\"</i>", body_style))
        story.append(Spacer(1, 10))
        
        story.append(Paragraph(f"<b>Confidence:</b> {int(data_content.get('confidence_score', 0) * 100)}%", body_style))
        story.append(Paragraph(f"<b>Source Document:</b> {data_content.get('doc_source', '')}", body_style))
        
    doc.build(story)
    
    # Save path in DB
    relative_path = f"/static/reports/{pdf_filename}"
    report.pdf_path = relative_path
    db.add(report)
    db.commit()
    
    return pdf_path

def build_docx_report(report_id: int, db: Session) -> str:
    report = db.query(Report).filter(Report.id == report_id).first()
    if not report:
        raise ValueError("Report not found")
        
    doc_meta = db.query(Document).filter(Document.id == report.document_id).first()
    company = db.query(Company).filter(Company.id == report.company_id).first()
    
    company_name = company.name if company else "N/A"
    doc_name = doc_meta.file_name if doc_meta else "N/A"
    
    docx_filename = f"report_{report_id}_{report.report_type}.docx"
    docx_path = os.path.join(settings.REPORT_DIR, docx_filename)
    
    doc = DocxDocument()
    
    # Style configuration
    styles = doc.styles
    normal_style = styles['Normal']
    normal_style.font.name = 'Arial'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(30, 41, 59)
    
    # Cover Section
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_run = title_p.add_run(report.title)
    title_run.font.size = Pt(26)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(15, 23, 42)
    
    sub_p = doc.add_paragraph()
    sub_run = sub_p.add_run("AI-Powered Executive Due Diligence Platform")
    sub_run.font.size = Pt(13)
    sub_run.font.color.rgb = RGBColor(100, 116, 139)
    
    doc.add_paragraph("\n" * 4)
    
    # Meta table
    table = doc.add_table(rows=2, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(3.2)
    table.columns[1].width = Inches(3.2)
    
    table.cell(0, 0).paragraphs[0].add_run(f"Company: {company_name}").bold = True
    table.cell(0, 1).paragraphs[0].add_run(f"Report Type: {report.report_type.upper()}").bold = True
    table.cell(1, 0).paragraphs[0].add_run(f"Document: {doc_name}")
    table.cell(1, 1).paragraphs[0].add_run(f"Date: {report.created_at.strftime('%Y-%m-%d')}")
    
    doc.add_page_break()
    
    # Report Data
    data_content = json.loads(report.content)
    
    if report.report_type == "risk":
        h1 = doc.add_heading(level=1)
        h1.add_run("Executive Summary").font.color.rgb = RGBColor(15, 23, 42)
        doc.add_paragraph(data_content.get("executive_summary", ""))
        
        h2 = doc.add_heading(level=1)
        h2.add_run("Overall Risk Score Assessment").font.color.rgb = RGBColor(15, 23, 42)
        doc.add_paragraph(f"Overall Risk Score: {data_content.get('overall_risk_score', 0)} / 100")
        doc.add_paragraph(f"Risk Severity Level: {data_content.get('risk_level', 'Unknown')}")
        
        h3 = doc.add_heading(level=2)
        h3.add_run("Risk Category Breakdown").font.color.rgb = RGBColor(59, 130, 246)
        breakdown = data_content.get("risk_breakdown", {})
        for k, v in breakdown.items():
            doc.add_paragraph(f"• {k.replace('_', ' ').title()}: {v}/100")
            
        h4 = doc.add_heading(level=1)
        h4.add_run("Key Risks Identified").font.color.rgb = RGBColor(15, 23, 42)
        for risk in data_content.get("key_risks", []):
            doc.add_paragraph(f"• {risk}")
            
        h5 = doc.add_heading(level=1)
        h5.add_run("Mitigation Recommendations").font.color.rgb = RGBColor(15, 23, 42)
        for rec in data_content.get("recommendations", []):
            doc.add_paragraph(f"• {rec}")
            
        h6 = doc.add_heading(level=1)
        h6.add_run("Final Consultant Opinion").font.color.rgb = RGBColor(15, 23, 42)
        doc.add_paragraph(data_content.get("final_consultant_opinion", ""))
        
    elif report.report_type == "investment":
        h1 = doc.add_heading(level=1)
        h1.add_run("Investment Summary").font.color.rgb = RGBColor(15, 23, 42)
        doc.add_paragraph(f"Investment Score: {data_content.get('investment_score', 0)} / 100")
        doc.add_paragraph(f"Recommendation Verdict: {data_content.get('recommendation', 'HOLD')}")
        doc.add_paragraph(f"Analyst Confidence Level: {data_content.get('confidence_score', 0)}%")
        
        h2 = doc.add_heading(level=1)
        h2.add_run("Strategic Strengths").font.color.rgb = RGBColor(15, 23, 42)
        for strength in data_content.get("strengths", []):
            doc.add_paragraph(f"• {strength}")
            
        h3 = doc.add_heading(level=1)
        h3.add_run("Key Weaknesses").font.color.rgb = RGBColor(15, 23, 42)
        for weakness in data_content.get("weaknesses", []):
            doc.add_paragraph(f"• {weakness}")
            
        h4 = doc.add_heading(level=1)
        h4.add_run("Valuation Insights").font.color.rgb = RGBColor(15, 23, 42)
        doc.add_paragraph(data_content.get("valuation_insight", ""))
        
        h5 = doc.add_heading(level=1)
        h5.add_run("Final Analyst Opinion").font.color.rgb = RGBColor(15, 23, 42)
        doc.add_paragraph(data_content.get("final_analyst_opinion", ""))
        
    else:  # Chat report
        h1 = doc.add_heading(level=1)
        h1.add_run("AI Document Intelligence Query").font.color.rgb = RGBColor(15, 23, 42)
        doc.add_paragraph(f"Question: {data_content.get('question', '')}").style.font.bold = True
        
        doc.add_paragraph("\nAnswer:")
        doc.add_paragraph(data_content.get("answer", ""))
        
        doc.add_paragraph("\nSupporting Document Evidence:")
        doc.add_paragraph(f"\"{data_content.get('supporting_evidence', '')}\"").style.font.italic = True
        
        doc.add_paragraph(f"\nConfidence: {int(data_content.get('confidence_score', 0) * 100)}%")
        doc.add_paragraph(f"Source Document: {data_content.get('doc_source', '')}")
        
    doc.save(docx_path)
    
    # Save path in DB
    relative_path = f"/static/reports/{docx_filename}"
    report.docx_path = relative_path
    db.add(report)
    db.commit()
    
    return docx_path
