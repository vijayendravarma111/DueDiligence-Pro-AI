import os
import logging
from pypdf import PdfReader
import docx

logger = logging.getLogger(__name__)


def extract_text_from_pdf(file_path: str) -> str:
    """Extract text from PDF using PyPDF."""
    text = ""
    try:
        reader = PdfReader(file_path)
        for page in reader.pages:
            extracted = page.extract_text()
            if extracted:
                text += extracted + "\n"
    except Exception as e:
        logger.error(f"Error extracting PDF text with PyPDF: {e}")
        raise ValueError(f"Failed to extract text from PDF: {str(e)}")
    return text.strip()


def extract_text_from_docx(file_path: str) -> str:
    """Extract text from DOCX using python-docx."""
    text_list = []
    try:
        doc = docx.Document(file_path)
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_list.append(paragraph.text.strip())
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                if row_text:
                    text_list.append(row_text)
    except Exception as e:
        logger.error(f"Error extracting DOCX text: {e}")
        raise ValueError(f"Failed to extract text from DOCX: {str(e)}")
    return "\n".join(text_list).strip()


def extract_text(file_path: str, file_type: str) -> str:
    """Main text extractor supporting PDF and DOCX."""
    ext = file_type.lower().strip(".")
    if ext == "pdf":
        return extract_text_from_pdf(file_path)
    elif ext in ["docx", "doc"]:
        return extract_text_from_docx(file_path)
    else:
        raise ValueError(f"Unsupported file format: {file_type}. Supported formats are PDF and DOCX.")


def chunk_text(text: str, chunk_size: int = 1000, chunk_overlap: int = 200) -> list[str]:
    """
    Splits document text into clean, overlapping character chunks.
    """
    if not text or not text.strip():
        return []

    chunks = []
    start = 0
    text_len = len(text)

    while start < text_len:
        end = min(start + chunk_size, text_len)
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)

        if end >= text_len:
            break

        start += chunk_size - chunk_overlap

    return chunks
