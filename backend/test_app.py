import os
import sys
import time
import unittest
from fastapi.testclient import TestClient
from docx import Document as DocxDocument
from pypdf import PdfWriter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from app.main import app

client = TestClient(app)


class TestDueDiligenceProAI(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        """Prepare sample PDF and DOCX files for testing."""
        cls.sample_pdf = "test_sample.pdf"
        cls.sample_docx = "test_sample.docx"

        # Create sample PDF using PyPDF
        writer = PdfWriter()
        writer.add_blank_page(width=612, height=792)
        with open(cls.sample_pdf, "wb") as f:
            writer.write(f)

        # Create sample DOCX using python-docx with rich due-diligence content
        doc = DocxDocument()
        doc.add_heading('Acme Corp Due Diligence Report 2026', 0)
        doc.add_paragraph('Company Overview: Acme Corp is a SaaS technology company with $15M annual recurring revenue.')
        doc.add_paragraph('Financial Performance: EBITDA margin is 25% with zero debt. Revenue grew 40% Year-over-Year.')
        doc.add_paragraph('Risk Assessment: Key risks include high customer concentration in top 3 clients.')
        doc.add_paragraph('Legal & Compliance: All software patents are registered. No active litigation.')
        doc.add_paragraph('Recommendations: Expand enterprise sales team and diversify customer base.')
        doc.save(cls.sample_docx)

    @classmethod
    def tearDownClass(cls):
        """Clean up test sample files."""
        for path in [cls.sample_pdf, cls.sample_docx]:
            if os.path.exists(path):
                try:
                    os.remove(path)
                except Exception:
                    pass

    def test_01_health_check(self):
        """Test health check endpoint and MySQL connection."""
        response = client.get("/health")
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data["status"], "healthy")
        self.assertEqual(data["database"], "ok")
        print("[OK] Health check passed (MySQL database is connected).")

    def test_02_user_registration_and_login(self):
        """Test user registration, login, JWT token issuance, and protected routes."""
        test_email = f"analyst_{int(time.time())}@example.com"
        password = "SecurePassword123"

        # 1. Register User
        reg_resp = client.post(
            "/api/v1/auth/register",
            json={"name": "Alice Analyst", "email": test_email, "password": password}
        )
        self.assertEqual(reg_resp.status_code, 201)
        reg_data = reg_resp.json()
        self.assertIn("access_token", reg_data)
        token = reg_data["access_token"]
        print(f"[OK] User registration passed for '{test_email}'.")

        # 2. Duplicate Registration Rejection
        dup_resp = client.post(
            "/api/v1/auth/register",
            json={"name": "Alice Analyst", "email": test_email, "password": password}
        )
        self.assertEqual(dup_resp.status_code, 400)
        print("[OK] Duplicate user registration rejected cleanly.")

        # 3. User Login
        login_resp = client.post(
            "/api/v1/auth/login",
            json={"email": test_email, "password": password}
        )
        self.assertEqual(login_resp.status_code, 200)
        login_data = login_resp.json()
        self.assertIn("access_token", login_data)
        print("[OK] User login passed.")

        # 4. Get Current User Profile
        headers = {"Authorization": f"Bearer {token}"}
        me_resp = client.get("/api/v1/auth/me", headers=headers)
        self.assertEqual(me_resp.status_code, 200)
        me_data = me_resp.json()
        self.assertEqual(me_data["email"], test_email)
        print("[OK] Protected /auth/me route passed.")

    def test_03_document_upload_and_processing(self):
        """Test DOCX document upload, text extraction, chunking, Chroma vector storage, and executive summary."""
        # 1. Register test user
        email = f"user_doc_{int(time.time())}@example.com"
        reg_resp = client.post(
            "/api/v1/auth/register",
            json={"name": "Bob Analyst", "email": email, "password": "Password123"}
        )
        token = reg_resp.json()["access_token"]
        headers = {"Authorization": f"Bearer {token}"}

        # 2. Upload DOCX Document
        with open(self.sample_docx, "rb") as f:
            upload_resp = client.post(
                "/api/v1/documents/upload",
                files={"file": ("Acme_Due_Diligence.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
                headers=headers
            )
        self.assertEqual(upload_resp.status_code, 201)
        doc_data = upload_resp.json()
        doc_id = doc_data["id"]
        self.assertEqual(doc_data["filename"], "Acme_Due_Diligence.docx")
        self.assertEqual(doc_data["status"], "completed")
        self.assertIsNotNone(doc_data.get("analysis"))
        print(f"[OK] DOCX upload & AI indexing passed for doc_id={doc_id}.")

        # 3. Retrieve Executive Summary
        summary_resp = client.get(f"/api/v1/documents/{doc_id}", headers=headers)
        self.assertEqual(summary_resp.status_code, 200)
        summary_data = summary_resp.json()["analysis"]["summary"]
        self.assertIn("overview", summary_data)
        self.assertIn("key_findings", summary_data)
        print("[OK] Executive Summary generation verified.")

        # 4. Perform RAG Question Answering
        ask_resp = client.post(
            f"/api/v1/documents/{doc_id}/ask",
            json={"question": "What is Acme Corp's annual revenue and EBITDA margin?"},
            headers=headers
        )
        self.assertEqual(ask_resp.status_code, 200)
        answer_data = ask_resp.json()
        self.assertIn("answer", answer_data)
        self.assertIsNotNone(answer_data.get("confidence_score"))
        print(f"[OK] RAG Q&A response: {answer_data['answer'][:100]}...")

        # 5. Retrieve Chat History
        history_resp = client.get(f"/api/v1/documents/{doc_id}/chat-history", headers=headers)
        self.assertEqual(history_resp.status_code, 200)
        history_data = history_resp.json()
        self.assertGreaterEqual(len(history_data), 1)
        print("[OK] Chat history retrieval verified.")

    def test_04_user_document_isolation(self):
        """Test multi-tenant security: Users can NEVER access other users' documents."""
        # User 1 uploads document
        u1_resp = client.post(
            "/api/v1/auth/register",
            json={"name": "User One", "email": f"u1_{int(time.time())}@example.com", "password": "Password123"}
        )
        u1_token = u1_resp.json()["access_token"]

        with open(self.sample_docx, "rb") as f:
            up_resp = client.post(
                "/api/v1/documents/upload",
                files={"file": ("Private_Doc.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
                headers={"Authorization": f"Bearer {u1_token}"}
            )
        u1_doc_id = up_resp.json()["id"]

        # User 2 attempts to access User 1's document
        u2_resp = client.post(
            "/api/v1/auth/register",
            json={"name": "User Two", "email": f"u2_{int(time.time())}@example.com", "password": "Password123"}
        )
        u2_token = u2_resp.json()["access_token"]
        u2_headers = {"Authorization": f"Bearer {u2_token}"}

        # User 2 tries to GET User 1's document details -> 404
        access_resp = client.get(f"/api/v1/documents/{u1_doc_id}", headers=u2_headers)
        self.assertEqual(access_resp.status_code, 404)

        # User 2 tries to ASK AI about User 1's document -> 404
        ask_resp = client.post(
            f"/api/v1/documents/{u1_doc_id}/ask",
            json={"question": "Tell me secrets"},
            headers=u2_headers
        )
        self.assertEqual(ask_resp.status_code, 404)
        print("[OK] User-specific document isolation successfully verified!")

    def test_05_invalid_file_rejection(self):
        """Test file type validation rejecting unsupported file extensions."""
        u_resp = client.post(
            "/api/v1/auth/register",
            json={"name": "Tester", "email": f"tester_{int(time.time())}@example.com", "password": "Password123"}
        )
        token = u_resp.json()["access_token"]

        invalid_resp = client.post(
            "/api/v1/documents/upload",
            files={"file": ("malicious_script.exe", b"binary content", "application/octet-stream")},
            headers={"Authorization": f"Bearer {token}"}
        )
        self.assertEqual(invalid_resp.status_code, 400)
        print("[OK] Invalid file type rejection passed.")


if __name__ == "__main__":
    unittest.main()
